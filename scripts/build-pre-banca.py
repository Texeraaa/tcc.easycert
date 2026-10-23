#!/usr/bin/env python3
"""Gera ~/Downloads/pre-banca.pptx a partir do template Muller e de docs/tcc/rascunhos/pre-banca.md.

Uso:
    python scripts/build-pre-banca.py              # gera o pptx
    python scripts/build-pre-banca.py --self-test  # roda testes unitários
    python scripts/build-pre-banca.py --inspect    # imprime estrutura do template
"""
from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = Path.home() / "Downloads" / "Muller_PreBanca_v2.pptx"
RASCUNHO_PATH = REPO_ROOT / "docs" / "tcc" / "rascunhos" / "pre-banca.md"
OUTPUT_PATH = Path.home() / "Downloads" / "pre-banca.pptx"
IMG_BASE_DIR = RASCUNHO_PATH.parent


# --- parser ---


@dataclass
class SlideContent:
    number: int
    heading: str                             # texto após "## Slide N — "
    bullets: list[str] = field(default_factory=list)
    paragraphs: list[str] = field(default_factory=list)
    blockquote: list[str] = field(default_factory=list)
    table_rows: list[list[str]] = field(default_factory=list)
    images: list[tuple[str, str]] = field(default_factory=list)  # (alt, resolved_absolute_path)
    has_falta: bool = False


_SLIDE_HEADER_RE = re.compile(r"^## Slide (\d{1,2})\s*(?:—|-)\s*(.*)$")
_IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def parse_rascunho(markdown: str) -> dict[int, SlideContent]:
    """Parse docs/tcc/rascunhos/pre-banca.md em dict {slide_num: SlideContent}.

    Regras documentadas na spec seção 2.9.
    """
    slides: dict[int, SlideContent] = {}
    current: SlideContent | None = None

    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        header = _SLIDE_HEADER_RE.match(line)
        if header:
            num = int(header.group(1))
            if not 1 <= num <= 14:
                raise ValueError(f"Slide fora da faixa 1-14: {num}")
            current = SlideContent(number=num, heading=header.group(2).strip())
            slides[num] = current
            continue
        if current is None:
            continue
        if line.startswith("## ") or line.startswith("# "):
            current = None
            continue
        if "[FALTA" in line:
            current.has_falta = True
        for alt, path in _IMAGE_RE.findall(line):
            resolved = _resolve_image_path(path)
            current.images.append((alt, resolved))
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells):
                continue
            current.table_rows.append(cells)
        elif stripped.startswith("> "):
            current.blockquote.append(stripped[2:].strip())
        elif stripped.startswith(">"):
            current.blockquote.append("")
        elif stripped.startswith("- "):
            without_img = _IMAGE_RE.sub("", stripped[2:]).strip()
            if without_img:
                current.bullets.append(without_img)
        else:
            current.paragraphs.append(stripped)

    return slides


def _resolve_image_path(path: str) -> str:
    p = Path(path)
    if not p.is_absolute():
        p = (IMG_BASE_DIR / path).resolve()
    return str(p)


_INLINE_MARK_RE = re.compile(r"(\*\*\*|\*\*|\*|_)")


def strip_inline_marks(text: str) -> str:
    """Remove **, *, _ do texto. Não interpreta — só remove."""
    return _INLINE_MARK_RE.sub("", text)


def normalize_for_diff(text: str) -> str:
    """Normaliza para o diff da seção 2.9: colapsa whitespace, remove bordas, strip inline marks."""
    text = strip_inline_marks(text)
    return re.sub(r"\s+", " ", text).strip()


def replace_text_preserving_runs(text_frame, new_text: str) -> None:
    """Substitui todo o texto de um TextFrame por new_text, preservando a formatação
    do PRIMEIRO run do PRIMEIRO parágrafo (fonte, cor, tamanho, bold, italic).

    Remove parágrafos extras do template. Se new_text tiver '\\n', vira múltiplos parágrafos.
    """
    if not text_frame.paragraphs:
        text_frame.text = new_text
        return
    first_para = text_frame.paragraphs[0]
    if not first_para.runs:
        first_para.add_run()
    template_run = first_para.runs[0]
    font_name = template_run.font.name
    font_size = template_run.font.size
    font_bold = template_run.font.bold
    font_italic = template_run.font.italic
    try:
        font_color_rgb = template_run.font.color.rgb
    except (AttributeError, TypeError):
        font_color_rgb = None

    tf_elem = text_frame._txBody
    for p in list(tf_elem.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}p"))[1:]:
        tf_elem.remove(p)
    for r in list(first_para._p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")):
        first_para._p.remove(r)

    lines = new_text.split("\n")
    first_run = first_para.add_run()
    first_run.text = lines[0]
    _apply_font(first_run, font_name, font_size, font_bold, font_italic, font_color_rgb)
    for extra in lines[1:]:
        p = text_frame.add_paragraph()
        run = p.add_run()
        run.text = extra
        _apply_font(run, font_name, font_size, font_bold, font_italic, font_color_rgb)


def _apply_font(run, name, size, bold, italic, color_rgb) -> None:
    if name is not None:
        run.font.name = name
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color_rgb is not None:
        from pptx.dml.color import RGBColor
        run.font.color.rgb = RGBColor(color_rgb[0], color_rgb[1], color_rgb[2])


# --- slide handlers ---


def find_title_and_body_frames(slide):
    """Localiza (title_tf, body_tf) de um slide.

    Title: shape placeholder cujo placeholder_format.idx == 0 (ou cujo nome contém 'title').
    Body (nesta ordem de preferência):
        1. Shape (placeholder OU text_box) cujo nome contenha 'conteúdo'/'conteudo'/'content'
           — identifica o 'Espaço Reservado para Conteúdo' do template mesmo quando foi
           convertido em TextBox.
        2. SlidePlaceholder não-título (idx != 0).
        3. Primeiro text_frame restante (fallback).

    Essa preferência existe porque alguns slides do template (ex.: 7, 8) têm uma TextBox de
    sub-título ('Objetivo Geral', 'Objetivos Específicos') ANTES do shape de corpo na ordem
    XML, e tanto o subtítulo quanto o corpo já não são placeholders.
    """
    title_tf = None
    body_content_tf = None       # shape nomeado como "Conteúdo"
    body_placeholder_tf = None   # placeholder não-título
    body_fallback_tf = None      # qualquer outro text_frame
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        name = (shape.name or "").lower()
        is_placeholder = shape.is_placeholder
        is_title = (
            "title" in name
            or (is_placeholder and shape.placeholder_format.idx == 0)
        )
        if is_title and title_tf is None:
            title_tf = shape.text_frame
            continue
        is_content_named = (
            "conteúdo" in name or "conteudo" in name or "content" in name
        )
        if is_content_named and body_content_tf is None:
            body_content_tf = shape.text_frame
        elif is_placeholder and body_placeholder_tf is None:
            body_placeholder_tf = shape.text_frame
        elif body_fallback_tf is None:
            body_fallback_tf = shape.text_frame
    body_tf = body_content_tf or body_placeholder_tf or body_fallback_tf
    return title_tf, body_tf


def process_generic_slide(slide, content: SlideContent) -> list[str]:
    """Preenche título (do heading do rascunho) e corpo (bullets ou blockquote ou paragraphs)."""
    logs: list[str] = []
    title_tf, body_tf = find_title_and_body_frames(slide)

    if title_tf is not None and content.heading:
        replace_text_preserving_runs(title_tf, content.heading)
        logs.append(f"título='{content.heading}'")

    if body_tf is not None:
        body_text = _render_body(content)
        if body_text:
            replace_text_preserving_runs(body_tf, body_text)
            logs.append(f"corpo com {body_text.count(chr(10)) + 1} linha(s)")
        else:
            logs.append("corpo vazio (nenhum conteúdo no rascunho)")
    return logs


def process_blockquote_slide(slide, content: SlideContent) -> list[str]:
    """Para slides onde o blockquote carrega título (1ª linha) + corpo (demais linhas).

    Aplica-se a slides 1, 2, 5, 7, 14. O heading do rascunho é organizacional e descartado
    quando há blockquote; o título visível vem da 1ª linha do blockquote.

    Slide 5 é especial: título fixo "Pergunta de pesquisa" (spec seção 2.3).

    Fallback: se o template não tem body_frame (slides 2 e 14), concatena o blockquote
    inteiro no title_frame como múltiplos parágrafos.
    """
    logs: list[str] = []
    title_tf, body_tf = find_title_and_body_frames(slide)

    quote = [strip_inline_marks(b) for b in content.blockquote if b]
    if not quote:
        logs.append("sem blockquote — fallback para handler genérico")
        return process_generic_slide(slide, content)

    if content.number == 5:
        title_text = "Pergunta de pesquisa"
    else:
        title_text = quote[0]
    body_text = "\n".join(quote[1:]) if len(quote) > 1 else ""

    # Slide 5 tem bullets adicionais no corpo além do blockquote.
    if content.bullets:
        body_bullets = "\n".join(strip_inline_marks(b) for b in content.bullets)
        body_text = f"{body_text}\n{body_bullets}" if body_text else body_bullets

    if body_tf is None and body_text:
        # Fallback: template sem body frame (slides 2, 14) — concatena no título.
        full_title = f"{title_text}\n{body_text}"
        if title_tf is not None:
            replace_text_preserving_runs(title_tf, full_title)
            logs.append(f"título multi-linha (sem body frame): {full_title!r}")
        else:
            logs.append("AVISO: nem title_tf nem body_tf — nada substituído")
        return logs

    if title_tf is not None:
        replace_text_preserving_runs(title_tf, title_text)
        logs.append(f"título='{title_text[:50]}{'…' if len(title_text) > 50 else ''}'")
    if body_tf is not None:
        if body_text:
            replace_text_preserving_runs(body_tf, body_text)
            logs.append(f"corpo com {body_text.count(chr(10)) + 1} linha(s)")
        else:
            logs.append("corpo vazio")
    return logs


def _render_body(content: SlideContent) -> str:
    """Transforma SlideContent em string multiline para o body do slide.

    Prioridade: bullets > blockquote > paragraphs. Nunca mistura — se há bullets, só bullets.
    """
    if content.bullets:
        return "\n".join(strip_inline_marks(b) for b in content.bullets)
    if content.blockquote:
        return "\n".join(strip_inline_marks(b) for b in content.blockquote if b)
    if content.paragraphs:
        return "\n".join(strip_inline_marks(p) for p in content.paragraphs)
    return ""


def check_dependencies() -> None:
    try:
        import pptx  # noqa: F401
    except ImportError:
        sys.exit(
            "Erro: biblioteca 'python-pptx' não encontrada.\n"
            "Instale com: pip install --user python-pptx"
        )


def check_files_exist() -> None:
    if not TEMPLATE_PATH.is_file():
        sys.exit(f"Erro: template não encontrado em {TEMPLATE_PATH}.\n"
                 f"Coloque o arquivo Muller_PreBanca_v2.pptx em ~/Downloads/.")
    if not RASCUNHO_PATH.is_file():
        sys.exit(f"Erro: rascunho não encontrado em {RASCUNHO_PATH}.")


def inspect_template() -> int:
    from pptx import Presentation

    prs = Presentation(str(TEMPLATE_PATH))
    for idx, slide in enumerate(prs.slides, start=1):
        print(f"\n=== Slide {idx} ===")
        for shape_idx, shape in enumerate(slide.shapes):
            kind = type(shape).__name__
            name = getattr(shape, "name", "?")
            print(f"  [{shape_idx}] {kind} name={name!r}")
            if shape.has_text_frame:
                for p_idx, para in enumerate(shape.text_frame.paragraphs):
                    text = "".join(r.text for r in para.runs)
                    print(f"      para[{p_idx}]: {text!r}")
            if shape.has_table:
                table = shape.table
                print(f"      table rows={len(table.rows)} cols={len(table.columns)}")
                for r_idx, row in enumerate(table.rows):
                    cells = [cell.text_frame.text for cell in row.cells]
                    print(f"      row[{r_idx}]: {cells!r}")
    return 0


def run_self_tests() -> int:
    tests = [
        _test_parse_rascunho_detecta_14_slides,
        _test_parse_rascunho_slide_3_tem_4_bullets,
        _test_parse_rascunho_slide_1_tem_blockquote,
        _test_parse_rascunho_slide_11_tem_tabela,
        _test_parse_rascunho_slide_10_tem_duas_imagens,
        _test_parse_rascunho_detecta_falta,
        _test_strip_inline_marks_bold,
        _test_strip_inline_marks_italic,
        _test_strip_inline_marks_combined,
        _test_normalize_for_diff_whitespace,
    ]
    failures: list[str] = []
    for test in tests:
        try:
            test()
            print(f"  OK  {test.__name__}")
        except AssertionError as exc:
            failures.append(f"FAIL {test.__name__}: {exc}")
            print(f"  FAIL {test.__name__}: {exc}")
    if failures:
        print(f"\n{len(failures)} teste(s) falharam.")
        return 1
    print(f"\nTodos os {len(tests)} teste(s) passaram.")
    return 0


def _test_parse_rascunho_detecta_14_slides() -> None:
    slides = parse_rascunho(RASCUNHO_PATH.read_text(encoding="utf-8"))
    assert len(slides) == 14, f"esperado 14 slides, obtido {len(slides)}"
    assert set(slides.keys()) == set(range(1, 15)), f"keys: {sorted(slides.keys())}"


def _test_parse_rascunho_slide_3_tem_4_bullets() -> None:
    slides = parse_rascunho(RASCUNHO_PATH.read_text(encoding="utf-8"))
    s3 = slides[3]
    assert len(s3.bullets) == 4, f"esperado 4 bullets no slide 3, obtido {len(s3.bullets)}"
    assert "atividades complementares" in s3.bullets[0]


def _test_parse_rascunho_slide_1_tem_blockquote() -> None:
    slides = parse_rascunho(RASCUNHO_PATH.read_text(encoding="utf-8"))
    s1 = slides[1]
    assert len(s1.blockquote) >= 3, f"blockquote do slide 1 deveria ter >=3 linhas: {s1.blockquote}"
    assert "Pedro Henrique Teixeira" in "\n".join(s1.blockquote)


def _test_parse_rascunho_slide_11_tem_tabela() -> None:
    slides = parse_rascunho(RASCUNHO_PATH.read_text(encoding="utf-8"))
    s11 = slides[11]
    assert len(s11.table_rows) >= 2, "tabela do slide 11 deveria ter header + linhas"
    header = s11.table_rows[0]
    assert header == ["Atividades", "Ago", "Set", "Out", "Nov"], f"header: {header}"
    sus_row = next((r for r in s11.table_rows if "SUS" in r[0]), None)
    assert sus_row is not None, "linha SUS não encontrada"
    assert sus_row[1] == "" and sus_row[2] == "" and sus_row[3] == "X" and sus_row[4] == "X"


def _test_parse_rascunho_slide_10_tem_duas_imagens() -> None:
    slides = parse_rascunho(RASCUNHO_PATH.read_text(encoding="utf-8"))
    s10 = slides[10]
    assert len(s10.images) == 2, f"esperado 2 imagens no slide 10: {s10.images}"
    paths = [p for _, p in s10.images]
    assert all("aluno-" in p for p in paths), f"paths inesperados: {paths}"
    assert all(p.endswith(".png") for p in paths)


def _test_parse_rascunho_detecta_falta() -> None:
    slides = parse_rascunho(RASCUNHO_PATH.read_text(encoding="utf-8"))
    s13 = slides[13]
    assert s13.has_falta, "slide 13 deveria ter has_falta=True (regulamento FMPFM)"


def _test_strip_inline_marks_bold() -> None:
    assert strip_inline_marks("**oi**") == "oi"
    assert strip_inline_marks("palavra **destacada** aqui") == "palavra destacada aqui"


def _test_strip_inline_marks_italic() -> None:
    assert strip_inline_marks("*oi*") == "oi"
    assert strip_inline_marks("_oi_") == "oi"


def _test_strip_inline_marks_combined() -> None:
    assert strip_inline_marks("***oi***") == "oi"
    assert strip_inline_marks("**_oi_**") == "oi"


def _test_normalize_for_diff_whitespace() -> None:
    assert normalize_for_diff("  hello   world  ") == "hello world"
    assert normalize_for_diff("a\tb") == "a b"


def main() -> int:
    check_dependencies()
    check_files_exist()

    from pptx import Presentation

    slides_content = parse_rascunho(RASCUNHO_PATH.read_text(encoding="utf-8"))
    prs = Presentation(str(TEMPLATE_PATH))

    GENERIC_SLIDES = {3, 4, 6, 8, 9, 12}
    BLOCKQUOTE_SLIDES = {1, 2, 5, 7, 14}

    for idx, slide in enumerate(prs.slides, start=1):
        content = slides_content.get(idx)
        if content is None:
            print(f"Slide {idx}: SEM conteúdo no rascunho — mantido como no template.")
            continue
        if idx in GENERIC_SLIDES:
            logs = process_generic_slide(slide, content)
        elif idx in BLOCKQUOTE_SLIDES:
            logs = process_blockquote_slide(slide, content)
        else:
            logs = ["handler específico pendente (slide 10, 11, 13)"]
        print(f"Slide {idx}: {'; '.join(logs) if logs else '(nada)'}")

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PATH))
    print(f"\nArquivo gerado: {OUTPUT_PATH}")
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--self-test", action="store_true", help="Roda testes unitários inline.")
    parser.add_argument("--inspect", action="store_true", help="Imprime estrutura do template.")
    args = parser.parse_args()
    if args.self_test:
        sys.exit(run_self_tests())
    if args.inspect:
        sys.exit(inspect_template())
    sys.exit(main())
